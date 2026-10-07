import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pymupdf
from reportlab.pdfgen import canvas
from fastapi.testclient import TestClient

from backend.app.services.ai_optimizer import optimize_resume
from backend.app.services.jd_analyzer import analyze_job
from backend.app.services.jobs import get_job_description, JobNotFoundError
from backend.app.services.keyword_extractor import (
    extract_candidate_phrases,
    extract_terms,
    normalize,
)
from backend.app.services.matcher import match_resume
from backend.app.services.resume_generator import render_docx, render_pdf
from backend.app.services.resume_parser import extract_text, parse_resume
from backend.app.services.quality_analyzer import analyze_resume_quality
from backend.app.services.resume_validation import validate_optimized_resume, ResumeValidationError
from backend.app.main import app
import backend.app.main as main_module


SAMPLE = """Alex Morgan
alex@example.com | +1 555 555 2222 | Mumbai, India
linkedin.com/in/alex-morgan
PROFESSIONAL SUMMARY
Finance graduate with financial analysis and Excel experience.
EDUCATION
Crestview University
Bachelor of Commerce in Finance
2022 - 2025
GPA: 3.8
PROFESSIONAL EXPERIENCE
Northstar Services
Finance Assistant
January 2024 - May 2025
• Reconciled monthly expense data.
INTERNSHIP EXPERIENCE
ABC Technologies
Data Analyst Intern
June 2025 - August 2025
• Analyzed sales data using Excel and SQL.
• Prepared reports for the finance team.
PROJECTS
Student Budget Tracker
Technologies: Excel, Python
• Built a budget tracker to review monthly spending.
SKILLS
Excel, SQL, Python, Data Analysis, Communication
CERTIFICATIONS
Bloomberg Market Concepts
ACHIEVEMENTS
Finance case competition finalist
LANGUAGES
English, Hindi
"""


class PipelineTests(unittest.TestCase):
    def test_safe_names_and_resume_ids(self):
        self.assertEqual(main_module._safe_upload_name(r"..\outside\My Resume.docx", ".docx"), "My Resume.docx")
        self.assertTrue(main_module._is_resume_id("a" * 32))
        self.assertFalse(main_module._is_resume_id("../outside"))

    def test_parse_docx_and_sections(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "resume.docx"
            doc = Document()
            doc.add_paragraph(SAMPLE)
            doc.save(path)
            parsed = parse_resume(extract_text(path))
            self.assertEqual(parsed["personal_info"]["name"], "Alex Morgan")
            self.assertEqual(parsed["personal_info"]["location"], "Mumbai, India")
            self.assertFalse(any("alex@example.com" in line for line in parsed["other"]))
            self.assertEqual(parsed["education"][0]["grade"], "GPA: 3.8")
            self.assertEqual(parsed["education"][0]["institution"], "Crestview University")
            self.assertEqual(parsed["experience"][0]["company"], "Northstar Services")
            self.assertEqual(parsed["experience"][0]["role"], "Finance Assistant")
            self.assertIn("SQL", parsed["skills"]["technical"])
            self.assertIn("Excel", parsed["skills"]["tools"])
            self.assertEqual(parsed["internships"][0]["company"], "ABC Technologies")
            self.assertEqual(parsed["internships"][0]["role"], "Data Analyst Intern")
            self.assertEqual(parsed["internships"][0]["start_date"], "June 2025")
            self.assertTrue(parsed["projects"])
            self.assertEqual(parsed["projects"][0]["name"], "Student Budget Tracker")
            self.assertIn("Excel", parsed["projects"][0]["technologies"])

    def test_reordered_sections_and_missing_contact(self):
        parsed = parse_resume("Taylor Candidate\nEDUCATION\nBA Economics, College\nPROJECTS\nMarket research case study\nSKILLS\nTableau, SQL")
        self.assertEqual(parsed["personal_info"]["email"], "")
        self.assertEqual(parsed["personal_info"]["phone"], "")
        self.assertIn("BA Economics, College", [parsed["education"][0]["degree"], parsed["education"][0]["institution"]])
        self.assertEqual(parsed["projects"][0]["name"], "Market research case study")
        self.assertEqual(parsed["experience"], [])

    def test_parse_pdf(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "resume.pdf"
            pdf = canvas.Canvas(str(path))
            pdf.drawString(70, 750, "Jordan Lee")
            pdf.drawString(70, 730, "EDUCATION")
            pdf.drawString(70, 710, "BBA, University")
            pdf.save()
            parsed = parse_resume(extract_text(path))
            self.assertEqual(parsed["personal_info"]["name"], "Jordan Lee")
            self.assertEqual(parsed["education"][0]["degree"], "BBA, University")

    def test_ujma_regression_sections_and_facts(self):
        source = """UJMA MULANI
EDUCATION
Master of Computer Application (MCA) | 2025–2027
School of Computer Studies, Sri Balaji University, Pune | GPA: 8.71
Bachelor of Computer Application (BCA) | 2022–2025
Alard College of Business Studies, Pune | GPA: 8.6
CERTIFICATIONS
MERN Full Stack Development — SmartBridge, NASSCOM FutureSkills Prime
SAP Certified Associate — Back-End Developer (ABAP Cloud) — In Progress, Expected 2026
ACHIEVEMENTS & LEADERSHIP
Academic Coordinator for Semester III, supporting the smooth conduct of academic activities.
1st Place — Chess, Intra-College Event
PROJECTS
SkillSphere — Intelligent Career Readiness & Placement Ecosystem
CodeX Arena — Real-Time Competitive Coding Battle Platform
Fuel Agency Management System"""
        parsed = parse_resume(source)
        self.assertEqual(len(parsed["education"]), 2)
        self.assertEqual(parsed["education"][0]["institution"], "School of Computer Studies, Sri Balaji University, Pune")
        self.assertEqual(parsed["education"][0]["gpa"], "GPA: 8.71")
        self.assertEqual(parsed["education"][0]["dates"], "2025–2027")
        self.assertEqual(parsed["education"][1]["institution"], "Alard College of Business Studies, Pune")
        self.assertEqual(parsed["education"][1]["gpa"], "GPA: 8.6")
        self.assertEqual(len(parsed["certifications"]), 2)
        self.assertEqual(len(parsed["leadership"]), 2)
        self.assertEqual(len(parsed["projects"]), 3)
        self.assertEqual(parsed["section_order"], ["education", "certifications", "leadership", "projects"])
        with patch("backend.app.services.ai_optimizer._ai_is_configured", return_value=False):
            optimized, _ = optimize_resume(parsed, {}, {}, {"matched": []})
        safe = validate_optimized_resume(parsed, optimized)
        self.assertEqual(safe["education"], parsed["education"])
        self.assertEqual(safe["certifications"], parsed["certifications"])
        self.assertEqual(safe["leadership"], parsed["leadership"])
        self.assertEqual(safe["projects"], parsed["projects"])
        with tempfile.TemporaryDirectory() as tmp:
            docx_path, pdf_path = Path(tmp) / "ujma.docx", Path(tmp) / "ujma.pdf"
            render_docx(safe, docx_path)
            render_pdf(safe, pdf_path)
            docx_text = " ".join(p.text for p in Document(docx_path).paragraphs)
            self.assertIn("Education", docx_text)
            self.assertIn("Certifications", docx_text)
            self.assertIn("Achievements & Leadership", docx_text)
            self.assertIn("SkillSphere", docx_text)
            self.assertIn("GPA: 8.71", docx_text)
            with pymupdf.open(pdf_path) as document:
                pdf_text = " ".join(page.get_text() for page in document)
            self.assertIn("Education", pdf_text)
            self.assertIn("Certifications", pdf_text)
            self.assertIn("Achievements & Leadership", pdf_text)
            self.assertIn("SkillSphere", pdf_text)

    def test_validation_rejects_missing_critical_section(self):
        original = parse_resume(SAMPLE)
        damaged = dict(original, education=[])
        with self.assertRaises(ResumeValidationError):
            validate_optimized_resume(original, damaged)

    def test_validation_rejects_incomplete_education_facts(self):
        original = parse_resume("Candidate Person\nEDUCATION\nMCA | 2025–2027\nABC University | CGPA: 9.2")
        optimized = dict(original, education=[{"degree": "MCA", "institution": "", "dates": "2025–2027", "gpa": ""}])
        with self.assertRaisesRegex(ResumeValidationError, "institution missing|GPA/CGPA missing"):
            validate_optimized_resume(original, optimized)

    def test_gemini_omission_is_rejected_and_source_education_restored(self):
        source = parse_resume("Candidate Person\nEDUCATION\nMCA | 2025–2027\nABC University | CGPA: 9.2")
        model_output = {"personal_info": source["personal_info"], "education": [
            {"degree": "MCA", "institution": "", "dates": "2025–2027", "gpa": ""}
        ]}
        with patch("backend.app.services.ai_optimizer._ai_is_configured", return_value=True), \
             patch("backend.app.services.ai_optimizer._optimize_with_gemini", return_value=model_output):
            optimized, _ = optimize_resume(source, {}, {}, {"matched": []})
        safe = validate_optimized_resume(source, optimized)
        self.assertEqual(safe["education"], source["education"])
        from backend.app.services.ai_optimizer import _build_system_prompt
        prompt = _build_system_prompt()
        self.assertIn("Do not return an incomplete Education entry", prompt)
        self.assertIn("Education institutions", prompt)
        self.assertIn("GPA/CGPA", prompt)

    def test_categorized_skills_survive_parse_optimization_matching_and_rendering(self):
        source = """Candidate Person
SKILLS
Data & Analysis:
Python, SQL, Data Management, Data Interpretation, Quantitative Problem Solving
Programming:
Java, Python, C++, JavaScript, SQL
Backend & APIs:
Spring Boot, FastAPI, REST APIs, Hibernate, Spring Data JPA, WebSockets, JWT Authentication
Databases:
MySQL, MongoDB Atlas
Frontend:
React.js, HTML5, CSS3, Thymeleaf, Bootstrap, Tailwind CSS
Core:
Object-Oriented Programming, Data Structures & Algorithms, MVC Architecture, Database Management, Authentication & Authorization
Tools:
Git, GitHub, Maven, Postman, Judge0 API, VS Code, Eclipse, Spring Tool Suite (STS)
Operating Systems:
Windows, Linux (Basic)"""
        parsed = parse_resume(source)
        expected = {
            "Data & Analysis": ["Python", "SQL", "Data Management", "Data Interpretation", "Quantitative Problem Solving"],
            "Programming": ["Java", "Python", "C++", "JavaScript", "SQL"],
            "Backend & APIs": ["Spring Boot", "FastAPI", "REST APIs", "Hibernate", "Spring Data JPA", "WebSockets", "JWT Authentication"],
            "Databases": ["MySQL", "MongoDB Atlas"],
            "Frontend": ["React.js", "HTML5", "CSS3", "Thymeleaf", "Bootstrap", "Tailwind CSS"],
            "Core": ["Object-Oriented Programming", "Data Structures & Algorithms", "MVC Architecture", "Database Management", "Authentication & Authorization"],
            "Tools": ["Git", "GitHub", "Maven", "Postman", "Judge0 API", "VS Code", "Eclipse", "Spring Tool Suite (STS)"],
            "Operating Systems": ["Windows", "Linux (Basic)"],
        }
        self.assertEqual(parsed["skills"], expected)
        self.assertIn("python", parsed["normalized_skills"])
        self.assertEqual(parsed["normalized_skills"].count("python"), 1)
        malicious = {"skills": {"other": ["Backend & APIs:", "Spring Boot", "Python"]}}
        with patch("backend.app.services.ai_optimizer._ai_is_configured", return_value=True), \
             patch("backend.app.services.ai_optimizer._optimize_with_gemini", return_value=malicious):
            optimized, _ = optimize_resume(parsed, {}, {}, {"matched": []})
        safe = validate_optimized_resume(parsed, optimized)
        self.assertEqual(safe["skills"], expected)
        self.assertEqual(safe["skills"]["Data & Analysis"][0], "Python")
        self.assertEqual(safe["skills"]["Programming"][1], "Python")
        self.assertFalse(any(value.endswith(":") for values in safe["skills"].values() for value in values))
        match = match_resume(safe, {"required_skills": ["Spring Boot", "Python"]})
        self.assertEqual(match["matched"], ["Spring Boot", "Python"])
        self.assertNotIn("backend apis", match["resume_skills_detected"])
        with tempfile.TemporaryDirectory() as tmp:
            docx_path, pdf_path = Path(tmp) / "skills.docx", Path(tmp) / "skills.pdf"
            render_docx(safe, docx_path)
            render_pdf(safe, pdf_path)
            docx_text = " ".join(p.text for p in Document(docx_path).paragraphs)
            with pymupdf.open(pdf_path) as pdf:
                pdf_text = " ".join(page.get_text() for page in pdf)
            for content in (docx_text, pdf_text):
                self.assertIn("Data & Analysis:", content)
                self.assertIn("Backend & APIs:", content)
                self.assertIn("Operating Systems:", content)
                self.assertIn("Spring Tool Suite (STS)", content)

    def test_ujma_docx_complete_api_pipeline_preserves_facts(self):
        import io
        source = """UJMA MULANI
EDUCATION
Master of Computer Application (MCA) | 2025–2027
School of Computer Studies, Sri Balaji University, Pune | GPA: 8.71
Bachelor of Computer Application (BCA) | 2022–2025
Alard College of Business Studies, Pune | GPA: 8.6
CERTIFICATIONS
MERN Full Stack Development — SmartBridge, NASSCOM FutureSkills Prime
SAP Certified Associate — Back-End Developer (ABAP Cloud) — In Progress, Expected 2026
ACHIEVEMENTS & LEADERSHIP
Academic Coordinator for Semester III, supporting the smooth conduct of academic activities.
1st Place — Chess, Intra-College Event
PROJECTS
SkillSphere — Intelligent Career Readiness & Placement Ecosystem
CodeX Arena — Real-Time Competitive Coding Battle Platform
Fuel Agency Management System"""
        document = Document()
        for line in source.splitlines():
            document.add_paragraph(line)
        content = io.BytesIO()
        document.save(content)
        client = TestClient(app)
        uploaded = client.post("/api/optimize-resume", data={"job_id": "finance-demo-001"}, files={
            "file": ("UJMA MULANI Resume.docx", content.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        })
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        sid = uploaded.json()["id"]
        with patch("backend.app.services.ai_optimizer._ai_is_configured", return_value=False):
            optimized = client.post(f"/api/resumes/{sid}/optimize")
        self.assertEqual(optimized.status_code, 200, optimized.text)
        education = optimized.json()["optimized_resume"]["education"]
        self.assertEqual(education[0]["institution"], "School of Computer Studies, Sri Balaji University, Pune")
        self.assertEqual(education[0]["gpa"], "GPA: 8.71")
        self.assertEqual(education[1]["institution"], "Alard College of Business Studies, Pune")
        self.assertEqual(education[1]["gpa"], "GPA: 8.6")
        generated = client.post(f"/api/resumes/{sid}/generate")
        self.assertEqual(generated.status_code, 200, generated.text)
        outputs = main_module.SESSIONS[sid]["output_paths"]
        self.assertEqual(outputs["docx"].name, "UJMA MULANI Resume_Optimized.docx")
        docx_text = " ".join(p.text for p in Document(outputs["docx"]).paragraphs)
        with pymupdf.open(outputs["pdf"]) as pdf:
            pdf_text = " ".join(page.get_text() for page in pdf)
        for content_text in (docx_text, pdf_text):
            self.assertIn("School of Computer Studies, Sri Balaji University, Pune", content_text)
            self.assertIn("GPA: 8.71", content_text)
            self.assertIn("Alard College of Business Studies, Pune", content_text)
            self.assertIn("GPA: 8.6", content_text)
            self.assertIn("MERN Full Stack Development", content_text)
            self.assertIn("Academic Coordinator for Semester III", content_text)
            self.assertIn("1st Place", content_text)
            self.assertIn("SkillSphere", content_text)
        for path in outputs.values(): path.unlink(missing_ok=True)
        main_module.SESSIONS.pop(sid, None)

    def test_jd_and_normalization(self):
        job = get_job_description("finance-demo-001")
        jd = analyze_job(job)
        self.assertIn("financial modeling", [normalize(term) for term in jd["required_skills"]])
        self.assertEqual(normalize("Microsoft Excel"), "excel")
        self.assertIn("financial analysis", extract_terms("Financial Analysis"))
        self.assertTrue(any(item["type"] == "tool" and item["normalized_name"] == "excel" for item in jd["requirements"]))
        self.assertTrue(any(item["importance"] == "preferred" for item in jd["requirements"]))
        modeling = next(item for item in jd["requirements"] if item["normalized_name"] == "financial modeling")
        self.assertEqual(modeling["importance"], "required")
        self.assertTrue(any(item["type"] == "responsibility" for item in jd["requirements"]))

    def test_nlp_candidate_and_importance(self):
        jd = analyze_job({
            "title": "Data Analyst",
            "description": "Requirements:\nExperience with risk assessment and dashboard development.\nPreferred Skills:\nNice to have customer segmentation."
        })
        names = {item["normalized_name"] for item in jd["requirements"]}
        self.assertIn("risk assessment", names)
        self.assertIn("dashboard development", names)
        customer_segmentation = next(item for item in jd["requirements"] if item["normalized_name"] == "customer segmentation")
        self.assertEqual(customer_segmentation["importance"], "preferred")

    def test_matching_and_missing(self):
        job = {"description": ""}
        jd = {"required_skills": ["excel", "sql", "power bi", "forecasting"]}
        resume = {"raw_text": "Excel SQL financial analysis", "skills": ["Excel", "SQL", "Financial Analysis"]}
        match = match_resume(resume, jd)
        self.assertIn("excel", match["matched"])
        self.assertIn("forecasting", match["related"])
        self.assertIn("power bi", match["missing"])

    def test_alias_related_and_no_false_power_bi_inference(self):
        resume = {"source_lines": ["MS Excel", "Data Analysis", "Excel dashboards"], "skills": {"technical": ["Data Analysis"], "tools": ["MS Excel"]}}
        jd = {"required_skills": ["Excel", "Financial Analysis", "Power BI"]}
        match = match_resume(resume, jd)
        outcomes = {item["requirement"]: item for item in match["items"]}
        self.assertEqual(outcomes["Excel"]["status"], "matched")
        self.assertEqual(outcomes["Excel"]["match_type"], "alias")
        self.assertEqual(outcomes["Financial Analysis"]["status"], "related")
        self.assertEqual(outcomes["Power BI"]["status"], "not_detected")

    def test_experience_details_survive_all_template_renderers(self):
        resume = {
            "personal_info": {"name": "Candidate"},
            "experience": [{
                "company": "ABC Ltd.", "role": "Software Developer", "location": "Pune",
                "start_date": "Jan 2026", "end_date": "Present",
                "description": "Worked on backend development.",
                "bullets": ["Developed REST APIs.", "Implemented database operations."],
                "technologies": ["Java", "Spring Boot", "MySQL"],
            }],
        }
        expected = ["ABC Ltd.", "Software Developer", "Pune", "Jan 2026", "Present",
                    "Worked on backend development.", "Developed REST APIs.",
                    "Implemented database operations.", "Java", "Spring Boot", "MySQL"]
        with tempfile.TemporaryDirectory() as tmp:
            for template in ("professional", "corporate", "finance", "fresher"):
                docx, pdf = Path(tmp) / f"{template}.docx", Path(tmp) / f"{template}.pdf"
                render_docx(resume, docx, template)
                render_pdf(resume, pdf, template)
                docx_text = " ".join(p.text for p in Document(docx).paragraphs)
                with pymupdf.open(pdf) as document:
                    pdf_text = " ".join(page.get_text() for page in document)
                for fact in expected:
                    self.assertIn(fact, docx_text, f"{template} DOCX omitted {fact}")
                    self.assertIn(fact, pdf_text, f"{template} PDF omitted {fact}")

    def test_unknown_sections_and_markdown_are_safely_normalized(self):
        resume = {
            "personal_info": {"name": "Candidate"},
            "section_order": ["summary", "research_experience", "training"],
            "section_heading_labels": {"research_experience": ["Research Experience"]},
            "summary": "**Analytical** candidate",
            "research_experience": [{"organization": "Lab", "details": "Built a model"}],
            "training": ["### Valuation", "__DCF basics__"],
        }
        with tempfile.TemporaryDirectory() as tmp:
            docx, pdf = Path(tmp) / "custom.docx", Path(tmp) / "custom.pdf"
            render_docx(resume, docx, "finance")
            render_pdf(resume, pdf, "finance")
            docx_text = " ".join(p.text for p in Document(docx).paragraphs)
            with pymupdf.open(pdf) as document:
                pdf_text = " ".join(page.get_text() for page in document)
            for content in ("Research Experience", "Organization: Lab", "Details: Built a model", "Valuation", "DCF basics"):
                self.assertNotIn(content.casefold(), docx_text.casefold())
                self.assertNotIn(content.casefold(), pdf_text.casefold())
            self.assertIn("Analytical", docx_text)
            self.assertIn("Analytical", pdf_text)
            self.assertNotIn("**", docx_text + pdf_text)
            self.assertNotIn("__", docx_text + pdf_text)

    def test_renderer_omits_parser_fields_and_normalizes_contact_and_education(self):
        resume = {
            "personal_info": {"name": "Candidate Person", "email": r"person\@example.com", "linkedin": "linkedin.com/in/person", "links": ["linkedin.com/in/person"]},
            "education": [{"institution": "Example University", "degree": "BCA", "dates": "2022 - 2025", "raw_text": "BCA Example University 2022 - 2025", "other": ["raw text: duplicate"]}],
            "skills": {"Sql": ["MySQL", "Python"], "soft_skills": ["Communication"]},
            "other": [],
            "debug": {"parser": "internal"},
        }
        with tempfile.TemporaryDirectory() as tmp:
            for template in ("professional", "corporate", "finance", "fresher"):
                docx, pdf = Path(tmp) / f"{template}.docx", Path(tmp) / f"{template}.pdf"
                render_docx(resume, docx, template)
                render_pdf(resume, pdf, template)
                docx_text = " ".join(p.text for p in Document(docx).paragraphs)
                with pymupdf.open(pdf) as document: pdf_text = " ".join(page.get_text() for page in document)
                for content in (docx_text, pdf_text):
                    self.assertIn("person@example.com", content)
                    self.assertIn("Example University", content)
                    self.assertIn("BCA", content)
                    self.assertIn("SQL", content)
                    self.assertNotIn("raw text:", content.lower())
                    self.assertNotIn("duplicate", content.lower())
                    self.assertNotIn("debug", content.lower())
                    self.assertNotIn("linkedin.com/in/person | linkedin.com/in/person", content)

    def test_markdown_ai_response_is_rejected_for_deterministic_fallback(self):
        source = parse_resume(SAMPLE)
        with patch("backend.app.services.ai_optimizer._ai_is_configured", return_value=True), \
             patch("backend.app.services.ai_optimizer._optimize_with_gemini", return_value="### Professional Summary\\n**Invented candidate**"):
            optimized, method = optimize_resume(source, {}, {}, {"matched": []})
        self.assertIn("Deterministic fallback", method)
        self.assertEqual(optimized["education"], source["education"])
        self.assertEqual(optimized["experience"], source["experience"])
        self.assertNotIn("Invented candidate", str(optimized))

    def test_output_validation_and_generation(self):
        source = parse_resume(SAMPLE)
        with patch("backend.app.services.ai_optimizer._ai_is_configured", return_value=False):
            optimized, method = optimize_resume(source, {}, {}, {"matched": ["Excel"]})
        self.assertIn("not integrated", method)
        self.assertNotIn("Power BI", [skill for values in optimized["skills"].values() for skill in values])
        self.assertEqual(optimized["internships"], source["internships"])
        with tempfile.TemporaryDirectory() as tmp:
            pro_docx = Path(tmp) / "professional.docx"
            corp_docx = Path(tmp) / "corporate.docx"
            pro_pdf = Path(tmp) / "professional.pdf"
            corp_pdf = Path(tmp) / "corporate.pdf"
            render_docx(optimized, pro_docx, "professional")
            render_docx(optimized, corp_docx, "corporate")
            render_pdf(optimized, pro_pdf, "professional")
            render_pdf(optimized, corp_pdf, "corporate")
            for output in (pro_docx, corp_docx, pro_pdf, corp_pdf):
                self.assertGreater(output.stat().st_size, 0)
            self.assertNotEqual(pro_docx.read_bytes(), corp_docx.read_bytes())
            self.assertNotEqual(pro_pdf.read_bytes(), corp_pdf.read_bytes())
            pro_document = Document(pro_docx)
            corp_document = Document(corp_docx)
            pro_paragraphs = [p.text for p in pro_document.paragraphs]
            corp_paragraphs = [p.text for p in corp_document.paragraphs]
            self.assertIn("ABC Technologies", " ".join(pro_paragraphs))
            self.assertIn("Data Analyst Intern", " ".join(pro_paragraphs))
            self.assertTrue(any(run.bold for p in Document(pro_docx).paragraphs if "ABC Technologies" in p.text for run in p.runs))
            self.assertNotEqual(pro_paragraphs, corp_paragraphs)
            self.assertEqual(pro_document.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertEqual(corp_document.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.LEFT)
            self.assertEqual(pro_document.paragraphs[0].runs[0].font.name, "Arial")
            self.assertEqual(corp_document.paragraphs[0].runs[0].font.name, "Times New Roman")
            with pymupdf.open(pro_pdf) as pro, pymupdf.open(corp_pdf) as corp:
                pro_text = " ".join(page.get_text() for page in pro).lower().split()
                corp_text = " ".join(page.get_text() for page in corp).lower().split()
                self.assertEqual(pro_text, corp_text)

    def test_evidence_and_quality_analysis(self):
        resume = parse_resume(SAMPLE)
        job_analysis = analyze_job(get_job_description("finance-demo-001"))
        match = match_resume(resume, job_analysis)
        excel = next(item for item in match["items"] if item["normalized_name"] == "excel")
        self.assertEqual(excel["status"], "matched")
        self.assertTrue(excel["evidence"])
        missing = next(item for item in match["items"] if item["normalized_name"] == "power bi")
        self.assertEqual(missing["status"], "not_detected")
        self.assertEqual(missing["evidence"], [])
        quality = analyze_resume_quality(resume, match)
        self.assertIn("jd_alignment", quality)
        self.assertEqual(quality["jd_alignment"]["matched"], match["counts"]["matched"])

    def test_api_end_to_end(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "input.docx"
            doc = Document()
            doc.add_paragraph(SAMPLE)
            doc.save(source)
            client = TestClient(app)
            self.assertEqual(client.get("/").status_code, 200)
            with source.open("rb") as stream:
                response = client.post(
                    "/api/optimize-resume",
                    data={"job_id": "finance-demo-001"},
                    files={"file": ("my resume.docx", stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
                )
            self.assertEqual(response.status_code, 200, response.text)
            payload = response.json()
            self.assertIn("missing", payload["match"])
            source_path = main_module.SESSIONS[payload["id"]]["source_path"]
            self.assertTrue(source_path.exists(), "Upload must remain available after analysis.")
            self.assertIn("my resume.docx", source_path.name)
            self.assertEqual(client.post(f"/api/resumes/{payload['id']}/generate").status_code, 409)
            with patch("backend.app.services.ai_optimizer._ai_is_configured", return_value=False):
                optimized = client.post(f"/api/resumes/{payload['id']}/optimize")
            self.assertEqual(optimized.status_code, 200, optimized.text)
            self.assertTrue(source_path.exists(), "Upload must remain available until output generation succeeds.")
            with patch("backend.app.main.render_docx", wraps=main_module.render_docx) as docx_renderer:
                generated = client.post(f"/api/resumes/{payload['id']}/generate", data={"template_id": "corporate"})
                docx_renderer.assert_called_once()
            self.assertEqual(generated.status_code, 200, generated.text)
            paths = main_module.SESSIONS[payload['id']]['output_paths']
            self.assertTrue(paths['docx'].is_file())
            self.assertTrue(paths['pdf'].is_file())
            self.assertEqual(paths['docx'].name, "my resume_Optimized.docx")
            self.assertEqual(paths['pdf'].name, "my resume_Optimized.pdf")
            self.assertFalse(source_path.exists(), "Successful output generation must remove only the original upload.")
            self.assertEqual(client.get(generated.json()["docx_url"]).status_code, 200)
            self.assertEqual(client.get(generated.json()["pdf_url"]).status_code, 200)
            paths['docx'].unlink(missing_ok=True)
            paths['pdf'].unlink(missing_ok=True)
            main_module.SESSIONS.pop(payload["id"], None)

    def test_failed_generation_retains_upload(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "input.docx"
            doc = Document()
            doc.add_paragraph(SAMPLE)
            doc.save(source)
            client = TestClient(app)
            with source.open("rb") as stream:
                response = client.post(
                    "/api/optimize-resume",
                    data={"job_id": "finance-demo-001"},
                    files={"file": ("retry me.docx", stream)}
                )
            self.assertEqual(response.status_code, 200, response.text)
            sid = response.json()["id"]
            source_path = main_module.SESSIONS[sid]["source_path"]
            with patch("backend.app.services.ai_optimizer._ai_is_configured", return_value=False):
                self.assertEqual(client.post(f"/api/resumes/{sid}/optimize").status_code, 200)
            with patch("backend.app.main.render_pdf", side_effect=OSError("simulated PDF failure")):
                failed = client.post(f"/api/resumes/{sid}/generate")
            self.assertEqual(failed.status_code, 500)
            self.assertTrue(source_path.exists(), "Failed generation must retain the original for debugging.")
            source_path.unlink(missing_ok=True)
            for path in main_module.OUTPUTS.glob("retry me_Optimized*"):
                path.unlink(missing_ok=True)
            main_module.SESSIONS.pop(sid, None)

    def test_batch_independent_results_filenames_and_limit(self):
        doc = Document()
        doc.add_paragraph(SAMPLE)
        import io
        payload = io.BytesIO()
        doc.save(payload)
        content = payload.getvalue()
        client = TestClient(app)
        with patch("backend.app.services.ai_optimizer._ai_is_configured", return_value=False):
            response = client.post("/api/optimize-resumes", data={"job_id": "finance-demo-001"}, files=[
                ("files", ("same resume.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
                ("files", ("same resume.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ])
        self.assertEqual(response.status_code, 200, response.text)
        items = response.json()["results"]
        self.assertEqual([item["status"] for item in items], ["completed", "completed"])
        self.assertEqual(items[0]["optimized_docx"], "same resume_Optimized.docx")
        self.assertEqual(items[1]["optimized_docx"], "same resume_Optimized_2.docx")
        self.assertTrue(all(Path(main_module.OUTPUTS / item["optimized_pdf"]).is_file() for item in items))
        for item in items:
            sid = item["id"]
            for path in main_module.SESSIONS[sid]["output_paths"].values(): path.unlink(missing_ok=True)
            main_module.SESSIONS.pop(sid, None)
        with patch("backend.app.services.ai_optimizer._ai_is_configured", return_value=False):
            partial = client.post("/api/optimize-resumes", data={"job_id": "finance-demo-001"}, files=[
                ("files", ("bad.docx", b"not a docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
                ("files", ("good.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ])
        partial_items = partial.json()["results"]
        self.assertEqual([item["status"] for item in partial_items], ["failed", "completed"])
        self.assertEqual(partial_items[0]["original_filename"], "bad.docx")
        for item in partial_items:
            if item["status"] == "completed":
                sid = item["id"]
                for path in main_module.SESSIONS[sid]["output_paths"].values(): path.unlink(missing_ok=True)
                main_module.SESSIONS.pop(sid, None)
        too_many = client.post("/api/optimize-resumes", data={"job_id": "finance-demo-001"}, files=[
            ("files", (f"{i}.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")) for i in range(6)
        ])
        self.assertEqual(too_many.status_code, 400)
        self.assertIn("Maximum 5 resumes", too_many.json()["detail"])

    def test_batch_five_mixed_pdf_docx_outputs(self):
        import io
        doc = Document()
        doc.add_paragraph(SAMPLE)
        docx_buffer = io.BytesIO()
        doc.save(docx_buffer)
        pdf_buffer = io.BytesIO()
        pdf = canvas.Canvas(pdf_buffer)
        for index, line in enumerate(SAMPLE.splitlines()):
            pdf.drawString(45, 780 - index * 18, line[:95])
        pdf.save()
        files = [("files", (f"Resume_{index}.docx", docx_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")) for index in range(1, 5)]
        files.append(("files", ("Data Analyst Resume.pdf", pdf_buffer.getvalue(), "application/pdf")))
        with patch("backend.app.services.ai_optimizer._ai_is_configured", return_value=False):
            response = TestClient(app).post("/api/optimize-resumes", data={"job_id": "finance-demo-001"}, files=files)
        self.assertEqual(response.status_code, 200, response.text)
        items = response.json()["results"]
        self.assertEqual(len(items), 5)
        self.assertTrue(all(item["status"] == "completed" for item in items))
        self.assertEqual(items[-1]["optimized_pdf"], "Data Analyst Resume_Optimized.pdf")
        for item in items:
            sid = item["id"]
            paths = main_module.SESSIONS[sid]["output_paths"]
            self.assertTrue(paths["docx"].is_file())
            self.assertTrue(paths["pdf"].is_file())
            for path in paths.values(): path.unlink(missing_ok=True)
            main_module.SESSIONS.pop(sid, None)

    # =========================================================================
    # PART 21 — REQUIRED EXPLICIT TESTS (TEST 1 to TEST 13)
    # =========================================================================

    def test_part21_01_job_retrieval(self):
        """TEST 1 — Job retrieval: Request demo job_id. Expected: Correct JD returned."""
        job = get_job_description("finance-demo-001")
        self.assertEqual(job["job_id"], "finance-demo-001")
        self.assertEqual(job["title"], "Financial Analyst")
        self.assertEqual(job["company"], "ABC Finance Ltd.")

        # Alias check
        demo_job = get_job_description("demo-001")
        self.assertEqual(demo_job["job_id"], "finance-demo-001")

    def test_part21_02_dynamic_job_id(self):
        """TEST 2 — Dynamic job ID: Create two different test jobs with different IDs.
        Verify that requesting job A returns job A and requesting job B returns job B.
        Do not return a hardcoded demo job."""
        job_a = get_job_description("finance-demo-001")
        job_b = get_job_description("tech-demo-002")

        self.assertEqual(job_a["job_id"], "finance-demo-001")
        self.assertEqual(job_a["title"], "Financial Analyst")

        self.assertEqual(job_b["job_id"], "tech-demo-002")
        self.assertEqual(job_b["title"], "Software Engineer")
        self.assertNotEqual(job_a["job_id"], job_b["job_id"])
        self.assertNotEqual(job_a["title"], job_b["title"])

    def test_part21_03_missing_job(self):
        """TEST 3 — Missing job: Request an unknown ID. Expected: 404-style error."""
        with self.assertRaises(JobNotFoundError):
            get_job_description("FINXL-UNKNOWN-999")

        client = TestClient(app)
        resp = client.get("/api/jobs/FINXL-UNKNOWN-999")
        self.assertEqual(resp.status_code, 404)
        data = resp.json()
        self.assertEqual(data.get("error"), "Job not found")
        self.assertEqual(data.get("job_id"), "FINXL-UNKNOWN-999")

    def test_part21_04_structured_jd_extraction(self):
        """TEST 4 — Structured JD extraction:
        Verify required_skills, preferred_skills, responsibilities, education are extracted."""
        job = get_job_description("finance-demo-001")
        jd = analyze_job(job)

        self.assertTrue(len(jd["required_skills"]) > 0)
        self.assertTrue(len(jd["preferred_skills"]) > 0)
        self.assertTrue(len(jd["responsibilities"]) > 0)
        self.assertTrue(len(jd["qualifications"]) > 0)

        # Confirm internal requirements list contains typed entries for all sections
        sources = {r["source"] for r in jd["requirements"]}
        self.assertIn("required_skills", sources)
        self.assertIn("preferred_skills", sources)
        self.assertIn("responsibilities", sources)
        self.assertIn("education", sources)

    def test_part21_05_required_vs_preferred(self):
        """TEST 5 — Required/preferred:
        Verify Financial Modeling = required, Python = preferred."""
        job = get_job_description("finance-demo-001")
        jd = analyze_job(job)

        modeling = next(r for r in jd["requirements"] if r["normalized_name"] == "financial modeling")
        self.assertEqual(modeling["importance"], "required")

        python_req = next(r for r in jd["requirements"] if r["normalized_name"] == "python")
        self.assertEqual(python_req["importance"], "preferred")

    def test_part21_06_nlp_separation(self):
        """TEST 6 — NLP separation:
        Verify 'financial planning, reporting, forecasting, and business analysis'
        does not become 'Financial PlanningForecasting'. Expected separate concepts."""
        text = "Support financial planning, reporting, forecasting, and business analysis activities."
        phrases = extract_candidate_phrases(text)
        phrase_names = [p["name"] for p in phrases]
        phrase_norms = [p["normalized_name"] for p in phrases]

        # Must NOT concatenate into Financial PlanningForecasting
        for name in phrase_names:
            self.assertNotIn("planningforecasting", name.lower().replace(" ", ""))

        # Must have distinct candidates
        self.assertTrue(any("financial planning" in norm or norm == "financial planning" for norm in phrase_norms))
        self.assertTrue(any("forecasting" in norm or norm == "forecasting" for norm in phrase_norms))
        self.assertTrue(any("business analysis" in norm or norm == "business analysis" for norm in phrase_norms))

    def test_part21_07_qualification_extraction(self):
        """TEST 7 — Qualification extraction: Verify degree and education fields are extracted."""
        job = get_job_description("finance-demo-001")
        jd = analyze_job(job)

        quals = jd["qualifications"]
        self.assertTrue(any("degree" in q.lower() or "bachelor" in q.lower() for q in quals))
        self.assertTrue(any("finance" in q.lower() for q in quals))
        self.assertTrue(any("accounting" in q.lower() for q in quals))
        self.assertTrue(any("computer science" in q.lower() for q in quals))

    def test_part21_08_responsibility_extraction(self):
        """TEST 8 — Responsibility extraction:
        Verify all 8 finance demo responsibilities are preserved individually."""
        job = get_job_description("finance-demo-001")
        jd = analyze_job(job)

        resps = jd["responsibilities"]
        self.assertEqual(len(resps), 8)
        self.assertEqual(resps[0], "Prepare financial reports and management reports.")
        self.assertEqual(resps[1], "Analyze revenue, expenses, and financial performance.")
        self.assertEqual(resps[7], "Assist with financial planning and business analysis.")

        resp_reqs = [r for r in jd["requirements"] if r["type"] == "responsibility"]
        self.assertEqual(len(resp_reqs), 8)

    def test_part21_09_normalization(self):
        """TEST 9 — Normalization:
        Verify Microsoft Excel → Excel, MS Excel → Excel,
        Financial modelling → Financial modeling, Accounting fundamentals → Accounting."""
        self.assertEqual(normalize("Microsoft Excel"), "excel")
        self.assertEqual(normalize("MS Excel"), "excel")
        self.assertEqual(normalize("Advanced Microsoft Excel"), "excel")
        self.assertEqual(normalize("Financial modelling"), "financial modeling")
        self.assertEqual(normalize("Accounting fundamentals"), "accounting")
        self.assertEqual(normalize("Data analytics"), "data analysis")

    def test_part21_10_matching(self):
        """TEST 10 — Matching: Verify matched, related, and not_detected states."""
        job = get_job_description("finance-demo-001")
        jd = analyze_job(job)

        resume = {
            "source_lines": [
                "Used Microsoft Excel and SQL to analyze sales data",
                "Prepared recurring reports to summarize revenue and operating expenses",
            ],
            "skills": {"tools": ["Microsoft Excel"], "technical": ["SQL"]}
        }
        match = match_resume(resume, jd)
        status_by_name = {item["requirement"]: item["status"] for item in match["items"]}

        # SQL is directly mentioned -> matched
        self.assertEqual(status_by_name["SQL"], "matched")
        # Excel is an alias of Microsoft Excel -> matched
        self.assertEqual(status_by_name["Excel"], "matched")
        # Financial Reporting is evidenced by recurring reports / operating expenses -> matched or related
        self.assertIn(status_by_name["Financial Reporting"], {"matched", "related"})
        # Power BI is not mentioned -> not_detected
        self.assertEqual(status_by_name["Power BI"], "not_detected")

    def test_part21_11_evidence(self):
        """TEST 11 — Evidence: Verify matched/related results include actual resume evidence."""
        job = get_job_description("finance-demo-001")
        jd = analyze_job(job)

        evidence_line = "Executed complex SQL queries to extract ledger records"
        resume = {
            "source_lines": [evidence_line],
            "skills": ["SQL"]
        }
        match = match_resume(resume, jd)
        sql_item = next(item for item in match["items"] if item["normalized_name"] == "sql")

        self.assertEqual(sql_item["status"], "matched")
        self.assertIn(evidence_line, sql_item["evidence"])

        power_bi_item = next(item for item in match["items"] if item["normalized_name"] == "power bi")
        self.assertEqual(power_bi_item["status"], "not_detected")
        self.assertEqual(power_bi_item["evidence"], [])

    def test_part21_12_no_hallucination(self):
        """TEST 12 — No hallucination: Use a resume that does NOT contain Power BI.
        Verify Power BI = not_detected, and verify Power BI is NOT added to the optimized resume."""
        resume = parse_resume(SAMPLE)
        # Verify source does not contain Power BI
        all_skills = [s for group in resume["skills"].values() for s in group]
        self.assertNotIn("Power BI", all_skills)

        job = get_job_description("finance-demo-001")
        jd = analyze_job(job)
        match = match_resume(resume, jd)

        power_bi_match = next(item for item in match["items"] if item["normalized_name"] == "power bi")
        self.assertEqual(power_bi_match["status"], "not_detected")

        # Run optimization (using deterministic or AI validation)
        with patch("backend.app.services.ai_optimizer._ai_is_configured", return_value=False):
            optimized, _ = optimize_resume(resume, job, jd, match)

        opt_skills = [s for group in optimized["skills"].values() for s in group]
        self.assertNotIn("Power BI", opt_skills)
        self.assertNotIn("power bi", [s.lower() for s in opt_skills])

    def test_part21_13_ai_failure_fallback(self):
        """TEST 13 — AI failure fallback: Simulate AI unavailable.
        Verify the system still returns deterministic source-preserving content."""
        resume = parse_resume(SAMPLE)
        job = get_job_description("finance-demo-001")
        jd = analyze_job(job)
        match = match_resume(resume, jd)

        with patch("backend.app.services.ai_optimizer._ai_is_configured", return_value=True), \
             patch("backend.app.services.ai_optimizer._optimize_with_gemini", side_effect=RuntimeError("AI API 503 Service Unavailable")):
            optimized, method = optimize_resume(resume, job, jd, match)

        self.assertIn("Deterministic fallback", method)
        self.assertIn("Alex Morgan", optimized["personal_info"]["name"])
        self.assertEqual(optimized["education"][0]["institution"], "Crestview University")


if __name__ == "__main__":
    unittest.main()
